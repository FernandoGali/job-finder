"""Monta o DOCX do currículo reescrito, num modelo que sistemas de triagem (ATS) leem bem.

Regras do modelo: uma coluna só; nada de tabelas, caixas de texto, imagens, cabeçalho ou rodapé
(muitos ATS ignoram o que está neles); títulos de seção padrão; fonte comum; marcadores simples.
"""

import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from jobfinder_api.resume_rewrite import ResumeRewrite

FONT = "Calibri"


def _heading(document: docx.Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)


def _bullet(document: docx.Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(0)


def build_docx(rewrite: ResumeRewrite) -> bytes:
    document = docx.Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)

    if rewrite.name:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(rewrite.name)
        run.bold = True
        run.font.size = Pt(18)
    if rewrite.headline:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(rewrite.headline).font.size = Pt(13)
    if rewrite.contacts:
        paragraph = document.add_paragraph(" | ".join(rewrite.contacts))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if rewrite.summary:
        _heading(document, "Resumo")
        document.add_paragraph(rewrite.summary)

    if rewrite.experiences:
        _heading(document, "Experiência")
        for experience in rewrite.experiences:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(0)
            title = " | ".join(part for part in (experience.role, experience.company) if part)
            paragraph.add_run(title).bold = True
            if experience.period:
                paragraph.add_run(f" | {experience.period}")
            for bullet in experience.bullets:
                _bullet(document, bullet)

    if rewrite.education:
        _heading(document, "Formação")
        for item in rewrite.education:
            _bullet(
                document, " | ".join(p for p in (item.course, item.institution, item.period) if p)
            )

    if rewrite.skills:
        _heading(document, "Competências")
        document.add_paragraph(", ".join(rewrite.skills))

    if rewrite.certifications:
        _heading(document, "Cursos e certificações")
        for certification in rewrite.certifications:
            _bullet(document, certification)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
