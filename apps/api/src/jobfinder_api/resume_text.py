"""Extrai o texto do currículo (PDF ou DOCX) para a análise."""

import io

import docx
import pypdf

from jobfinder_api.resumes import DOCX, PDF

# Um currículo real tem poucas páginas. Acima disso, recusamos em vez de cortar o texto.
MAX_CHARS = 100_000


class UnreadableResumeError(Exception):
    """O arquivo não tem texto que dê para ler (ex.: PDF escaneado como imagem)."""


def extract_text(content: bytes, content_type: str) -> str:
    try:
        if content_type == PDF:
            text = _pdf_text(content)
        elif content_type == DOCX:
            text = _docx_text(content)
        else:
            raise UnreadableResumeError(f"tipo não suportado: {content_type}")
    except UnreadableResumeError:
        raise
    except Exception as exc:  # arquivo corrompido: qualquer erro das bibliotecas de leitura
        raise UnreadableResumeError(str(exc)) from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        raise UnreadableResumeError("sem texto")
    if len(text) > MAX_CHARS:
        raise UnreadableResumeError("texto longo demais")
    return text


def _pdf_text(content: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # Muitos currículos usam tabelas para organizar as seções.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
