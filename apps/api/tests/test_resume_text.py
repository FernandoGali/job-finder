import io

import docx
import pypdf
import pytest

from jobfinder_api.resume_text import MAX_CHARS, UnreadableResumeError, extract_text
from jobfinder_api.resumes import DOCX, PDF
from tests.conftest import pdf_with_text


def docx_with(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_le_texto_do_pdf():
    text = extract_text(pdf_with_text("Analista de Dados", "Python e SQL"), PDF)

    assert text == "Analista de Dados\nPython e SQL"


def test_le_paragrafos_e_tabelas_do_docx():
    content = docx_with(["Maria Silva", "", "Analista de Dados"], table=[["Python", "SQL"]])

    text = extract_text(content, DOCX)

    assert text.splitlines() == ["Maria Silva", "Analista de Dados", "Python", "SQL"]


def test_pdf_sem_texto_como_um_escaneado_e_recusado():
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(UnreadableResumeError, match="sem texto"):
        extract_text(buffer.getvalue(), PDF)


def test_arquivo_corrompido_e_recusado():
    with pytest.raises(UnreadableResumeError):
        extract_text(b"%PDF-1.7 isto nao e um pdf de verdade", PDF)
    with pytest.raises(UnreadableResumeError):
        extract_text(b"PK\x03\x04 zip quebrado", DOCX)


def test_texto_longo_demais_e_recusado_em_vez_de_cortado():
    content = docx_with(["palavra " * (MAX_CHARS // 8 + 10)])

    with pytest.raises(UnreadableResumeError, match="longo demais"):
        extract_text(content, DOCX)


def test_tipo_desconhecido_e_recusado():
    with pytest.raises(UnreadableResumeError, match="tipo não suportado"):
        extract_text(b"qualquer", "image/png")
