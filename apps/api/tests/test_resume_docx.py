import io

import docx

from jobfinder_api.resume_docx import build_docx
from jobfinder_api.resume_rewrite import Education, Experience, ResumeRewrite

REWRITE = ResumeRewrite(
    name="Maria Silva",
    contacts=["maria@exemplo.com", "São Paulo, SP"],
    headline="Analista de Dados",
    summary="Analista com foco em dados.",
    experiences=[
        Experience(
            company="Empresa X",
            role="Analista de BI",
            period="2021 a 2025",
            bullets=["Automatizei 12 relatórios semanais.", "Estruturei dashboards."],
        ),
        Experience(company="Empresa Y", role=None, period=None, bullets=[]),
    ],
    education=[Education(institution="Universidade Y", course="Engenharia", period=None)],
    skills=["SQL", "Python"],
    certifications=["Linux Unhatched"],
    fixes=["Ajuste que não vai para o arquivo."],
    questions=["Pergunta que não vai para o arquivo?"],
)


def _open(content: bytes) -> docx.Document:
    return docx.Document(io.BytesIO(content))


def test_docx_tem_uma_coluna_secoes_padrao_e_o_conteudo():
    document = _open(build_docx(REWRITE))

    lines = [p.text for p in document.paragraphs]
    assert lines[:3] == ["Maria Silva", "Analista de Dados", "maria@exemplo.com | São Paulo, SP"]
    for title in ("RESUMO", "EXPERIÊNCIA", "FORMAÇÃO", "COMPETÊNCIAS", "CURSOS E CERTIFICAÇÕES"):
        assert title in lines
    assert "Analista de BI | Empresa X | 2021 a 2025" in lines
    assert "Automatizei 12 relatórios semanais." in lines
    assert "Empresa Y" in lines
    assert "Engenharia | Universidade Y" in lines
    assert "SQL, Python" in lines
    # O que é só para a pessoa (ajustes e perguntas) não entra no currículo.
    assert not any("não vai para o arquivo" in line for line in lines)
    # Nada que os ATS costumam ignorar: tabelas, imagens, cabeçalho ou rodapé.
    assert document.tables == []
    assert document.inline_shapes._inline_lst == []
    assert all(not s.header.paragraphs[0].text for s in document.sections)
    assert document.styles["Normal"].font.name == "Calibri"


def test_docx_sem_secoes_vazias():
    empty = ResumeRewrite(
        headline=None,
        summary="",
        experiences=[],
        education=[],
        skills=[],
        certifications=[],
        questions=[],
    )

    lines = [p.text for p in _open(build_docx(empty)).paragraphs]

    assert lines == []
